"""
DragWheel conversion engine.

Everything here runs locally. No network calls.
Backends:
  - Pillow (+ pillow-heif)          images
  - ffmpeg (bundled by imageio-ffmpeg) video / audio
  - PyMuPDF                         PDF, EPUB, SVG, HTML/Markdown -> PDF
  - openpyxl                        spreadsheets
  - Microsoft Office (COM) or LibreOffice, if installed   DOCX/PPTX/XLS(X) -> PDF etc.
  - stdlib + py7zr                  archives (zip/7z/tar…), subtitles; RAR via 7-Zip/WinRAR/tar.exe
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tarfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import logging
logging.getLogger().setLevel(logging.WARNING)   # pdf2docx is chatty

IS_WIN = sys.platform == "win32"
OUTPUT_DIR: str | None = None   # None = save next to the original
NO_WINDOW = 0x08000000 if IS_WIN else 0  # CREATE_NO_WINDOW

# ----------------------------------------------------------------------------
# Categories
# ----------------------------------------------------------------------------
EXT_CATEGORY = {
    # images
    **{e: "image" for e in "jpg jpeg png webp heic heif tif tiff avif bmp gif ico svg psd tga dds jp2 j2k jpx pcx ppm pgm pbm qoi icns".split()},
    # video
    **{e: "video" for e in "mp4 mov mkv webm avi wmv m4v flv mpg mpeg 3gp ts mts m2ts vob ogv mxf f4v asf".split()},
    # audio
    **{e: "audio" for e in "mp3 m4a wav flac ogg oga opus aiff aif wma aac ac3 mka caf m4b amr au dts".split()},
    "pdf": "pdf",
    **{e: "document" for e in "docx doc rtf odt md markdown txt html htm".split()},
    **{e: "slides" for e in "pptx ppt odp key".split()},
    **{e: "spreadsheet" for e in "xlsx xlsm xls ods csv tsv".split()},
    **{e: "subtitle" for e in "srt vtt ass ssa".split()},
    **{e: "ebook" for e in "epub mobi fb2 cbz xps oxps".split()},
    **{e: "archive" for e in "zip tar gz tgz bz2 tbz2 xz txz 7z rar iso cab".split()},
}

CATEGORY_LABEL = {
    "image": "Image", "video": "Video", "audio": "Audio", "pdf": "PDF",
    "document": "Document", "slides": "Slides", "spreadsheet": "Spreadsheet",
    "subtitle": "Subtitles", "ebook": "E-book", "archive": "Archive",
    "folder": "Folder", "mixed": "Files", "other": "File",
}


def ext_of(p: str | Path) -> str:
    name = Path(p).name.lower()
    for suf, e in ((".tar.gz", "tgz"), (".tar.bz2", "tbz2"), (".tar.xz", "txz")):
        if name.endswith(suf):
            return e
    e = Path(p).suffix.lower().lstrip(".")
    return {"tbz": "tbz2"}.get(e, e)


def category_of(paths: list[str]) -> str:
    cats = set()
    for p in paths:
        if os.path.isdir(p):
            cats.add("folder")
        else:
            cats.add(EXT_CATEGORY.get(ext_of(p), "other"))
    if len(cats) == 1:
        return cats.pop()
    return "mixed"


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
class ConversionError(Exception):
    pass


# ----------------------------------------------------------------------------
# Cancellation: each job runs with a CancelToken (thread-local). Cancel kills any
# external process of the job, loops call check_cancel(), and outputs are deleted.
# ----------------------------------------------------------------------------
import threading

_ctx = threading.local()


class Cancelled(Exception):
    pass


class CancelToken:
    def __init__(self):
        self.cancelled = False
        self.procs: list[subprocess.Popen] = []
        self.paths: list[Path] = []        # every output path handed out for this job
        self.on_path = None                # callback(kind, path) — lets a parent process clean up
        self._lock = threading.Lock()

    def cancel(self):
        self.cancelled = True
        with self._lock:
            for p in self.procs:
                try:
                    if p.poll() is None:
                        p.kill()
                except Exception:
                    pass


def _token() -> CancelToken | None:
    return getattr(_ctx, "token", None)


def check_cancel():
    t = _token()
    if t and t.cancelled:
        raise Cancelled()


def _remember(path: Path, kind: str = "out") -> Path:
    t = _token()
    if t is not None:
        if kind == "out":
            t.paths.append(path)
        if t.on_path:
            try:
                t.on_path(kind, str(path))
            except Exception:
                pass
    return path


def temp_dir() -> Path:
    """A private temp folder for the current job (reported so it can be removed if the job is killed)."""
    import tempfile
    return _remember(Path(tempfile.mkdtemp(prefix="dragwheel-")), "tmp")


def unique_path(path: Path) -> Path:
    if not path.exists():
        return _remember(path)
    name = path.name
    suf = ""
    for multi in (".tar.gz", ".tar.bz2", ".tar.xz"):
        if name.lower().endswith(multi):
            suf = name[-len(multi):]
            break
    else:
        suf = path.suffix
    stem = name[: len(name) - len(suf)] if suf else name
    i = 2
    while True:
        cand = path.with_name(f"{stem} ({i}){suf}")
        if not cand.exists():
            return _remember(cand)
        i += 1


def _target_dir(src: Path) -> Path:
    if OUTPUT_DIR and os.path.isdir(OUTPUT_DIR):
        return Path(OUTPUT_DIR)
    return src.parent


def out_for(src: str | Path, ext: str, tag: str = "") -> Path:
    src = Path(src)
    stem = src.name[: -len(".tar.gz")] if src.name.lower().endswith(".tar.gz") else src.stem
    name = f"{stem}{(' ' + tag) if tag else ''}.{ext}"
    return unique_path(_target_dir(src) / name)


def out_dir_for(src: str | Path, tag: str) -> Path:
    src = Path(src)
    return unique_path(_target_dir(src) / f"{src.stem} {tag}")


def run(cmd: list[str], cwd: str | None = None, ok_codes=(0,)) -> None:
    """Run an external tool; killed immediately if the job is cancelled."""
    check_cancel()
    import tempfile
    t = _token()
    with tempfile.TemporaryFile() as err:
        proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=err, stdin=subprocess.DEVNULL,
                                cwd=cwd, creationflags=NO_WINDOW)
        if t is not None:
            with t._lock:
                t.procs.append(proc)
        try:
            while True:
                try:
                    proc.wait(timeout=0.2)
                    break
                except subprocess.TimeoutExpired:
                    if t is not None and t.cancelled:
                        proc.kill()
                        proc.wait()
                        raise Cancelled()
        finally:
            if t is not None:
                with t._lock:
                    if proc in t.procs:
                        t.procs.remove(proc)
        if t is not None and t.cancelled:
            raise Cancelled()
        if proc.returncode not in ok_codes:
            err.seek(0)
            tail = err.read().decode("utf-8", "replace").strip().splitlines()[-4:]
            raise ConversionError("\n".join(tail) or f"{Path(cmd[0]).name} failed")


def parse_time(t: str) -> str:
    """Accept '90', '1:30', '00:01:30.5' -> ffmpeg time string."""
    t = str(t).strip()
    if not re.fullmatch(r"[\d:.]+", t):
        raise ConversionError(f"Invalid time: {t}")
    return t


# ----------------------------------------------------------------------------
# ffmpeg
# ----------------------------------------------------------------------------
_FFMPEG = None


def ffmpeg() -> str:
    global _FFMPEG
    if _FFMPEG:
        return _FFMPEG
    exe = shutil.which("ffmpeg")
    if not exe:
        try:
            import imageio_ffmpeg
            exe = imageio_ffmpeg.get_ffmpeg_exe()
        except Exception:
            exe = None
    if not exe:
        raise ConversionError("ffmpeg not found. Run: pip install imageio-ffmpeg")
    _FFMPEG = exe
    return exe


def cover_art(src) -> bytes | None:
    """The album cover embedded in a music file (MP3, M4A, FLAC, ...) as PNG bytes, or None."""
    try:
        r = subprocess.run([ffmpeg(), "-hide_banner", "-loglevel", "error", "-i", str(src), "-an",
                            "-map", "0:v:0?", "-frames:v", "1", "-f", "image2pipe", "-vcodec", "png", "-"],
                           capture_output=True, timeout=8, stdin=subprocess.DEVNULL, creationflags=NO_WINDOW)
        return r.stdout if r.returncode == 0 and len(r.stdout) > 100 else None
    except Exception:
        return None


def ff(src, out, *args, pre: tuple = ()) -> Path:
    run([ffmpeg(), "-hide_banner", "-loglevel", "error", "-y", *pre, "-i", str(src), *args, str(out)])
    return Path(out)


VIDEO_ARGS = {
    "mp4": ["-c:v", "libx264", "-preset", "medium", "-crf", "21", "-pix_fmt", "yuv420p",
            "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart"],
    "mov": ["-c:v", "libx264", "-crf", "20", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "192k"],
    "mkv": ["-c:v", "libx264", "-crf", "20", "-c:a", "aac", "-b:a", "192k"],
    "webm": ["-c:v", "libvpx-vp9", "-b:v", "0", "-crf", "33", "-deadline", "good", "-cpu-used", "5",
             "-row-mt", "1", "-c:a", "libopus", "-b:a", "128k"],
    "avi": ["-c:v", "mpeg4", "-q:v", "4", "-c:a", "libmp3lame", "-q:a", "3"],
    "wmv": ["-c:v", "wmv2", "-q:v", "4", "-c:a", "wmav2", "-b:a", "192k"],
}

AUDIO_ARGS = {
    "mp3": ["-vn", "-c:a", "libmp3lame", "-q:a", "2"],
    "m4a": ["-vn", "-c:a", "aac", "-b:a", "192k"],
    "wav": ["-vn", "-c:a", "pcm_s16le"],
    "flac": ["-vn", "-c:a", "flac"],
    "ogg": ["-vn", "-c:a", "libvorbis", "-q:a", "5"],
    "opus": ["-vn", "-c:a", "libopus", "-b:a", "128k"],
    "aiff": ["-vn", "-c:a", "pcm_s16be"],
    "wma": ["-vn", "-c:a", "wmav2", "-b:a", "192k"],
    "ac3": ["-vn", "-c:a", "ac3", "-b:a", "448k"],
}


def video_to(fmt):
    def f(src, p):
        return ff(src, out_for(src, fmt), *VIDEO_ARGS[fmt])
    return f


def audio_to(fmt):
    def f(src, p):
        try:
            return ff(src, out_for(src, fmt), *AUDIO_ARGS[fmt])
        except ConversionError as e:
            if "does not contain any stream" in str(e) or "matches no streams" in str(e):
                raise ConversionError("This file has no audio track") from None
            raise
    return f


def video_to_gif(src, p):
    out = out_for(src, "gif")
    vf = "fps=12,scale='min(640,iw)':-1:flags=lanczos,split[a][b];[a]palettegen=stats_mode=diff[p];[b][p]paletteuse=dither=bayer:bayer_scale=4"
    return ff(src, out, "-vf", vf, "-loop", "0")


def video_gif_clip(src, p):
    """Make GIF… for videos: the part between Start and End, 12 fps, up to 640 px wide, good colours."""
    def secs(t):
        parts = [float(x) for x in parse_time(t).split(":")]
        return sum(v * 60 ** i for i, v in enumerate(reversed(parts)))
    start, end = secs(p.get("Start", "0")), secs(p.get("End", "5"))
    if end <= start:
        raise ConversionError("End must be after Start")
    out = out_for(src, "gif", "clip")
    vf = "fps=12,scale='min(640,iw)':-1:flags=lanczos,split[a][b];[a]palettegen=stats_mode=diff[p];[b][p]paletteuse=dither=bayer:bayer_scale=4"
    return ff(src, out, "-t", f"{end - start:.3f}", "-vf", vf, "-loop", "0", pre=("-ss", f"{start:.3f}"))


def video_to_hevc(src, p):
    """H.265 in MP4: about half the size of H.264 at the same quality (plays on Windows 10/11, phones)."""
    return ff(src, out_for(src, "mp4", "hevc"), "-c:v", "libx265", "-preset", "medium", "-crf", "26",
              "-pix_fmt", "yuv420p", "-tag:v", "hvc1", "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart",
              "-x265-params", "log-level=error")


def video_compress(src, p):
    ext = "mp4"
    return ff(src, out_for(src, ext, "compressed"), "-c:v", "libx264", "-preset", "slow", "-crf", "28",
              "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "96k", "-movflags", "+faststart")


def video_scale(h):
    def f(src, p):
        return ff(src, out_for(src, "mp4", f"{h}p"), "-vf", f"scale=-2:'min({h},ih)'",
                  "-c:v", "libx264", "-crf", "23", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "128k")
    return f


def video_mute(src, p):
    return ff(src, out_for(src, ext_of(src), "muted"), "-c:v", "copy", "-an")


def media_trim(src, p):
    start, end = parse_time(p["Start"]), parse_time(p["End"])
    e = _norm(ext_of(src))
    if EXT_CATEGORY.get(ext_of(src)) == "video":
        ext = e if e in VIDEO_ARGS else "mp4"
        args = VIDEO_ARGS[ext]
    else:
        ext = e if e in AUDIO_ARGS else "mp3"
        args = AUDIO_ARGS[ext]
    return ff(src, out_for(src, ext, "trimmed"), "-ss", start, "-to", end, *args)


def video_rotate(src, p):
    return ff(src, out_for(src, "mp4", "rotated"), "-vf", "transpose=1",
              "-c:v", "libx264", "-crf", "20", "-pix_fmt", "yuv420p", "-c:a", "copy")


def video_snapshot(src, p):
    return ff(src, out_for(src, "png", "frame"), "-frames:v", "1", pre=("-ss", parse_time(p["At time"])))


def video_speed2x(src, p):
    return ff(src, out_for(src, "mp4", "2x"), "-filter_complex", "[0:v]setpts=0.5*PTS[v];[0:a]atempo=2.0[a]",
              "-map", "[v]", "-map", "[a]?", "-c:v", "libx264", "-crf", "22", "-pix_fmt", "yuv420p")


def audio_compress(src, p):
    return ff(src, out_for(src, "mp3", "compressed"), "-vn", "-c:a", "libmp3lame", "-b:a", "96k")


def audio_normalize(src, p):
    e = ext_of(src)
    args = AUDIO_ARGS.get(e, AUDIO_ARGS["mp3"])
    return ff(src, out_for(src, e if e in AUDIO_ARGS else "mp3", "normalized"), "-af", "loudnorm=I=-16:TP=-1.5:LRA=11", *args)


def audio_mono(src, p):
    e = ext_of(src)
    args = AUDIO_ARGS.get(e, AUDIO_ARGS["mp3"])
    return ff(src, out_for(src, e if e in AUDIO_ARGS else "mp3", "mono"), "-ac", "1", *args)


def audio_fade(src, p):
    e = ext_of(src)
    args = AUDIO_ARGS.get(e, AUDIO_ARGS["mp3"])
    return ff(src, out_for(src, e if e in AUDIO_ARGS else "mp3", "faded"),
              "-af", "afade=t=in:d=2,areverse,afade=t=in:d=2,areverse", *args)


def media_merge(paths, p):
    """Concatenate media files (re-encodes for safety)."""
    first = Path(paths[0])
    is_video = EXT_CATEGORY.get(ext_of(first)) == "video"
    ext = "mp4" if is_video else ext_of(first)
    out = out_for(first, ext, "merged")
    inputs = []
    for x in paths:
        inputs += ["-i", str(x)]
    n = len(paths)
    if is_video:
        streams = "".join(f"[{i}:v:0][{i}:a:0]" for i in range(n))
        fc = f"{streams}concat=n={n}:v=1:a=1[v][a]"
        maps = ["-map", "[v]", "-map", "[a]", "-c:v", "libx264", "-crf", "21", "-pix_fmt", "yuv420p", "-c:a", "aac"]
    else:
        streams = "".join(f"[{i}:a:0]" for i in range(n))
        fc = f"{streams}concat=n={n}:v=0:a=1[a]"
        maps = ["-map", "[a]", *AUDIO_ARGS.get(ext, AUDIO_ARGS["mp3"])[1:]]
    run([ffmpeg(), "-hide_banner", "-loglevel", "error", "-y", *inputs, "-filter_complex", fc, *maps, str(out)])
    return out


# ----------------------------------------------------------------------------
# Images
# ----------------------------------------------------------------------------
def _pil():
    from PIL import Image, ImageOps
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except Exception:
        pass
    return Image, ImageOps


def open_image(src):
    Image, ImageOps = _pil()
    if ext_of(src) == "svg":
        import pymupdf
        doc = pymupdf.open(str(src))
        pix = doc[0].get_pixmap(dpi=192, alpha=True)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        img.load()
        return img
    img = Image.open(str(src))
    try:
        img = ImageOps.exif_transpose(img)
    except Exception:
        pass
    return img


def _flatten(img, bg=(255, 255, 255)):
    Image, _ = _pil()
    if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
        img = img.convert("RGBA")
        base = Image.new("RGB", img.size, bg)
        base.paste(img, mask=img.split()[-1])
        return base
    return img.convert("RGB")


SAVE_KW = {
    "jpg": dict(format="JPEG", quality=92, optimize=True),
    "png": dict(format="PNG", optimize=True),
    "webp": dict(format="WEBP", quality=90, method=6),
    "tiff": dict(format="TIFF", compression="tiff_lzw"),
    "bmp": dict(format="BMP"),
    "gif": dict(format="GIF"),
    "avif": dict(format="AVIF", quality=75),
    "heic": dict(format="HEIF", quality=85),
    "ico": dict(format="ICO", sizes=[(16, 16), (24, 24), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)]),
    "pdf": dict(format="PDF", resolution=150),
    "tga": dict(format="TGA", compression="tga_rle"),
    "jp2": dict(format="JPEG2000", quality_mode="dB", quality_layers=[38]),
    "icns": dict(format="ICNS"),
}


def save_image(img, out: Path, fmt: str, **override):
    kw = {**SAVE_KW[fmt], **override}
    if fmt in ("jpg", "bmp", "pdf", "jp2") or (fmt == "heic" and img.mode not in ("RGB", "RGBA")):
        img = _flatten(img)
    elif fmt == "gif":
        img = img.convert("RGBA") if img.mode not in ("P", "L") else img
    elif img.mode not in ("RGB", "RGBA", "L", "LA"):
        img = img.convert("RGBA")
    if fmt in ("ico", "icns"):
        side = max(img.size)
        Image, _ = _pil()
        sq = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        sq.paste(img.convert("RGBA"), ((side - img.width) // 2, (side - img.height) // 2))
        img = sq
        if fmt == "icns" and side > 1024:
            img = img.resize((1024, 1024), Image.LANCZOS)
    img.save(str(out), **kw)
    return out


def image_to(fmt):
    def f(src, p):
        img = open_image(src)
        if getattr(img, "is_animated", False) and fmt in ("webp", "gif"):
            out = out_for(src, fmt)
            img.save(str(out), save_all=True, **{k: v for k, v in SAVE_KW[fmt].items()})
            return out
        return save_image(img, out_for(src, fmt), fmt)
    return f


def image_tool(tag, transform, keep_ext=True, fmt_override=None):
    def f(src, p):
        img = open_image(src)
        img = transform(img, p)
        e = ext_of(src)
        fmt = fmt_override or ("jpg" if e == "jpeg" else "tiff" if e == "tif" else "heic" if e == "heif" else e)
        if fmt == "svg":
            fmt = "png"
        return save_image(img, out_for(src, fmt, tag), fmt)
    return f


def _resize(img, p):
    pct = float(str(p["Scale %"]).strip("% ")) / 100.0
    if not 0 < pct <= 10:
        raise ConversionError("Scale must be between 1 and 1000%")
    Image, _ = _pil()
    w, h = max(1, round(img.width * pct)), max(1, round(img.height * pct))
    return img.resize((w, h), Image.LANCZOS)


def _square(img, p):
    s = min(img.size)
    l, t = (img.width - s) // 2, (img.height - s) // 2
    return img.crop((l, t, l + s, t + s))


def image_crop_resize(src, p):
    """p: rotate (0/90/180/270, clockwise), flip_h, flip_v, box (l,t,r,b) in the rotated image,
    size (w,h) of the output, format ('' = same as source)."""
    Image, _ = _pil()
    img = open_image(src)
    rot = int(p.get("rotate", 0)) % 360
    if rot:
        img = img.rotate(-rot, expand=True)
    if p.get("flip_h"):
        img = img.transpose(Image.FLIP_LEFT_RIGHT)
    if p.get("flip_v"):
        img = img.transpose(Image.FLIP_TOP_BOTTOM)
    box = p.get("box")
    if box:
        l, t, r, b = (int(round(v)) for v in box)
        l, t = max(0, l), max(0, t)
        r, b = min(img.width, r), min(img.height, b)
        if r - l >= 1 and b - t >= 1 and (l, t, r, b) != (0, 0, img.width, img.height):
            img = img.crop((l, t, r, b))
    size = p.get("size")
    if size:
        w, h = max(1, int(size[0])), max(1, int(size[1]))
        if (w, h) != img.size:
            img = img.resize((w, h), Image.LANCZOS)
    e = ext_of(src)
    fmt = p.get("format") or ("jpg" if e == "jpeg" else "tiff" if e == "tif" else "heic" if e == "heif" else e)
    if fmt not in SAVE_KW or fmt == "pdf":
        fmt = "png"
    return save_image(img, out_for(src, fmt, "edited"), fmt)


def _strip(img, p):
    Image, _ = _pil()
    clean = Image.new(img.mode, img.size)
    clean.putdata(list(img.getdata())) if img.mode == "P" else clean.paste(img)
    if img.mode == "P":
        clean.putpalette(img.getpalette())
    return clean


def image_compress(src, p):
    img = open_image(src)
    e = ext_of(src)
    has_alpha = img.mode in ("RGBA", "LA", "P")
    if e in ("png", "gif", "bmp", "tif", "tiff") and has_alpha:
        # palette-quantize keeps transparency, usually 50-70% smaller
        q = img.convert("RGBA").quantize(colors=256, method=2)
        return save_image(q, out_for(src, "png", "compressed"), "png")
    if e == "webp":
        return save_image(img, out_for(src, "webp", "compressed"), "webp", quality=72)
    Image, _ = _pil()
    if max(img.size) > 3000:
        img.thumbnail((3000, 3000), Image.LANCZOS)
    return save_image(img, out_for(src, "jpg", "compressed"), "jpg", quality=72, progressive=True)


# ----------------------------------------------------------------------------
# Image extras (1.2): resize presets, upscale, sharpen, watermark, stitch, GIF,
# remove background (AI), blur faces, copy text (OCR)
# ----------------------------------------------------------------------------
def _asset(*parts) -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base.joinpath("assets", *parts)


RESIZE_PRESETS = ["50%", "25%", "1920 px (1080p)", "3840 px (4K)", "1080 × 1080", "1080 × 1350",
                  "< 1 MB", "< 500 KB"]


def _same_fmt(src):
    e = ext_of(src)
    fmt = "jpg" if e == "jpeg" else "tiff" if e == "tif" else "heic" if e == "heif" else e
    return "png" if fmt not in SAVE_KW or fmt in ("pdf", "ico", "icns") else fmt


def _center_crop(img, ratio):
    w, h = img.size
    if w / h > ratio:
        nw = int(h * ratio)
        return img.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
    nh = int(w / ratio)
    return img.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))


def _under_size(img, src, limit):
    """Smallest change that gets the file under `limit` bytes: lower quality first, then smaller size."""
    Image, _ = _pil()
    alpha = img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)
    fmt = "webp" if alpha else "jpg"
    work = img.convert("RGBA") if alpha else _flatten(img)
    for _ in range(12):
        for q in (90, 82, 74, 66, 58, 50, 42):
            check_cancel()
            buf = io.BytesIO()
            if fmt == "jpg":
                work.save(buf, format="JPEG", quality=q, optimize=True, progressive=True)
            else:
                work.save(buf, format="WEBP", quality=q, method=4)
            if buf.tell() <= limit:
                out = out_for(src, fmt, "resized")
                out.write_bytes(buf.getvalue())
                return out
        work = work.resize((max(1, int(work.width * 0.8)), max(1, int(work.height * 0.8))), Image.LANCZOS)
    raise ConversionError("Couldn't make the file small enough")


def image_resize_preset(src, p):
    Image, _ = _pil()
    img = open_image(src)
    choice = p.get("Output size", "50%")
    if choice.startswith("<"):
        kb = 1024 if "MB" in choice else 500
        return _under_size(img, src, kb * 1024)
    if choice.endswith("%"):
        f = float(choice.rstrip("%")) / 100
        img = img.resize((max(1, round(img.width * f)), max(1, round(img.height * f))), Image.LANCZOS)
    elif "×" in choice:
        w, h = (int(x) for x in choice.replace(" ", "").split("×"))
        img = _center_crop(img, w / h).resize((w, h), Image.LANCZOS)
    else:                                                    # "1920 px": longest side, never enlarged
        side = int(choice.split()[0])
        if max(img.size) > side:
            img.thumbnail((side, side), Image.LANCZOS)
    fmt = _same_fmt(src)
    return save_image(img, out_for(src, fmt, "resized"), fmt)


def image_upscale(src, p):
    from PIL import ImageFilter
    Image, _ = _pil()
    img = open_image(src)
    if img.mode not in ("RGB", "RGBA", "L", "LA"):
        img = img.convert("RGBA")
    if img.width * img.height > 40_000_000:
        raise ConversionError("This image is already very large")
    img = img.resize((img.width * 2, img.height * 2), Image.LANCZOS)
    img = img.filter(ImageFilter.UnsharpMask(radius=2, percent=60, threshold=2))
    fmt = _same_fmt(src)
    return save_image(img, out_for(src, fmt, "2x"), fmt)


def image_sharpen(src, p):
    from PIL import ImageFilter
    img = open_image(src)
    if img.mode not in ("RGB", "RGBA", "L", "LA"):
        img = img.convert("RGBA")
    img = img.filter(ImageFilter.UnsharpMask(radius=2, percent=120, threshold=3))
    fmt = _same_fmt(src)
    return save_image(img, out_for(src, fmt, "sharpened"), fmt)


def _font(size):
    from PIL import ImageFont
    windir = os.environ.get("WINDIR", r"C:\Windows")
    for name in ("segoeuib.ttf", "segoeui.ttf", "arialbd.ttf", "arial.ttf", "DejaVuSans-Bold.ttf"):
        for folder in (Path(windir) / "Fonts", Path("/usr/share/fonts/truetype/dejavu")):
            f = folder / name
            if f.exists():
                return ImageFont.truetype(str(f), size)
    return ImageFont.load_default()


WATERMARK_POSITIONS = ["Bottom right", "Bottom left", "Center", "Tiled"]


def image_watermark(src, p):
    from PIL import ImageDraw
    Image, _ = _pil()
    text = (p.get("Text") or "").strip()
    if not text:
        raise ConversionError("Type the watermark text")
    pos = p.get("Position", "Bottom right")
    img = open_image(src).convert("RGBA")
    W, H = img.size
    size = max(12, int(min(W, H) * (0.05 if pos != "Tiled" else 0.045)))
    font = _font(size)
    layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    l, t, r, b = d.textbbox((0, 0), text, font=font)
    tw, th = r - l, b - t
    if pos == "Tiled":
        tile = Image.new("RGBA", (tw + size * 3, th + size * 3), (0, 0, 0, 0))
        ImageDraw.Draw(tile).text((size, size), text, font=font, fill=(255, 255, 255, 95),
                                  stroke_width=max(1, size // 18), stroke_fill=(0, 0, 0, 50))
        tile = tile.rotate(30, expand=True)
        for y in range(-tile.height // 2, H, tile.height):
            for x in range(-tile.width // 2, W, tile.width):
                layer.paste(tile, (x, y), tile)
    else:
        m = int(min(W, H) * 0.03)
        if pos == "Center":
            xy = ((W - tw) // 2 - l, (H - th) // 2 - t)
        elif pos == "Bottom left":
            xy = (m - l, H - th - m - t)
        else:
            xy = (W - tw - m - l, H - th - m - t)
        d.text(xy, text, font=font, fill=(255, 255, 255, 170), stroke_width=max(1, size // 16), stroke_fill=(0, 0, 0, 110))
    img.alpha_composite(layer)
    fmt = _same_fmt(src)
    return save_image(img, out_for(src, fmt, "watermarked"), fmt)


def images_stitch(direction):
    def f(paths, p):
        Image, _ = _pil()
        imgs = [open_image(x).convert("RGBA") for x in paths]
        if direction == "h":
            h = min(i.height for i in imgs)
            imgs = [i.resize((max(1, round(i.width * h / i.height)), h), Image.LANCZOS) for i in imgs]
            canvas = Image.new("RGBA", (sum(i.width for i in imgs), h), (0, 0, 0, 0))
            x = 0
            for i in imgs:
                canvas.paste(i, (x, 0)); x += i.width
        else:
            w = min(i.width for i in imgs)
            imgs = [i.resize((w, max(1, round(i.height * w / i.width))), Image.LANCZOS) for i in imgs]
            canvas = Image.new("RGBA", (w, sum(i.height for i in imgs)), (0, 0, 0, 0))
            y = 0
            for i in imgs:
                canvas.paste(i, (0, y)); y += i.height
        alpha = any(i.getextrema()[3][0] < 255 for i in imgs)
        fmt = "png" if alpha else "jpg"
        return save_image(canvas, out_for(paths[0], fmt, "stitched"), fmt)
    return f


def images_to_gif(paths, p):
    Image, _ = _pil()
    try:
        secs = float(str(p.get("Seconds per image", "0.5")).replace(",", "."))
    except ValueError:
        raise ConversionError("Seconds per image must be a number")
    secs = min(30.0, max(0.02, secs))
    first = open_image(paths[0])
    box = first.size
    if max(box) > 1200:
        f = 1200 / max(box)
        box = (round(box[0] * f), round(box[1] * f))
    frames = []
    for x in paths:
        check_cancel()
        im = _flatten(open_image(x))
        im.thumbnail(box, Image.LANCZOS)
        frame = Image.new("RGB", box, (255, 255, 255))
        frame.paste(im, ((box[0] - im.width) // 2, (box[1] - im.height) // 2))
        frames.append(frame)
    out = out_for(paths[0], "gif", "animated")
    frames[0].save(str(out), save_all=True, append_images=frames[1:], duration=int(secs * 1000), loop=0, optimize=True)
    return out


# ---- AI models (downloaded once, then offline)
MODELS = {
    "isnet": {"file": "isnet-general-use.onnx", "size": 178648008,
              "sha256": "60920e99c45464f2ba57bee2ad08c919a52bbf852739e96947fbb4358c0d964a",
              "url": "https://github.com/danielgatis/rembg/releases/download/v0.0.0/isnet-general-use.onnx"},
}


def models_dir() -> Path:
    base = os.environ.get("LOCALAPPDATA") or str(Path.home() / ".local" / "share")
    d = Path(base) / "DragWheel" / "models"
    d.mkdir(parents=True, exist_ok=True)
    return d


def model_path(name) -> Path:
    return models_dir() / MODELS[name]["file"]


def model_ready(name) -> bool:
    f = model_path(name)
    return f.is_file() and f.stat().st_size == MODELS[name]["size"]


_SESSIONS = {}


def _onnx(name):
    if name not in _SESSIONS:
        if not model_ready(name):
            raise ConversionError("The AI model isn't downloaded yet")
        import onnxruntime as ort
        opts = ort.SessionOptions()
        opts.log_severity_level = 3
        _SESSIONS[name] = ort.InferenceSession(str(model_path(name)), opts, providers=["CPUExecutionProvider"])
    return _SESSIONS[name]


def image_remove_bg(src, p):
    import numpy as np
    Image, _ = _pil()
    img = open_image(src)
    rgb = _flatten(img)
    sess = _onnx("isnet")
    check_cancel()
    x = np.asarray(rgb.resize((1024, 1024), Image.LANCZOS), dtype=np.float32)
    x = x / max(float(x.max()), 1e-6) - 0.5
    x = x.transpose(2, 0, 1)[None]
    pred = sess.run(None, {sess.get_inputs()[0].name: x})[0][0, 0]
    lo, hi = float(pred.min()), float(pred.max())
    pred = (pred - lo) / max(hi - lo, 1e-6)
    mask = Image.fromarray((pred * 255).astype("uint8")).resize(rgb.size, Image.LANCZOS)
    out_img = img.convert("RGBA")
    if img.mode in ("RGBA", "LA"):                      # keep existing transparency too
        from PIL import ImageChops
        mask = ImageChops.multiply(mask, out_img.split()[-1])
    out_img.putalpha(mask)
    bbox = mask.point(lambda v: 255 if v > 8 else 0).getbbox()
    if not bbox:
        raise ConversionError("No subject found in this image")
    return save_image(out_img, out_for(src, "png", "no background"), "png")


def image_blur_faces(src, p):
    import cv2
    import numpy as np
    from PIL import ImageDraw, ImageFilter
    Image, _ = _pil()
    img = open_image(src)
    rgb = _flatten(img)
    scale = min(1.0, 1600 / max(rgb.size))
    small = rgb.resize((max(1, int(rgb.width * scale)), max(1, int(rgb.height * scale))), Image.LANCZOS) if scale < 1 else rgb
    bgr = cv2.cvtColor(np.asarray(small), cv2.COLOR_RGB2BGR)
    det = cv2.FaceDetectorYN.create(str(_asset("models", "face_detection_yunet_2023mar.onnx")), "",
                                    (bgr.shape[1], bgr.shape[0]), 0.6, 0.3, 5000)
    _, faces = det.detect(bgr)
    if faces is None or len(faces) == 0:
        raise ConversionError("No faces found in this image")
    out = img.convert("RGBA") if img.mode in ("RGBA", "LA") else rgb.copy()
    for f in faces:
        x, y, w, h = (float(v) / scale for v in f[:4])
        pad = 0.18
        box = (int(max(0, x - w * pad)), int(max(0, y - h * pad)),
               int(min(out.width, x + w * (1 + pad))), int(min(out.height, y + h * (1 + pad))))
        region = out.crop(box)
        radius = max(6, int(max(region.size) / 6))
        blurred = region.filter(ImageFilter.GaussianBlur(radius))
        mask = Image.new("L", region.size, 0)
        ImageDraw.Draw(mask).ellipse([0, 0, region.width - 1, region.height - 1], fill=255)
        mask = mask.filter(ImageFilter.GaussianBlur(max(2, radius // 3)))
        out.paste(blurred, box[:2], mask)
    fmt = _same_fmt(src)
    return save_image(out, out_for(src, fmt, "blurred"), fmt)


OCR_PS = r"""
param([string]$Path)
$ErrorActionPreference = 'Stop'
[Console]::OutputEncoding = [Text.Encoding]::UTF8
Add-Type -AssemblyName System.Runtime.WindowsRuntime
$null = [Windows.Storage.StorageFile, Windows.Storage, ContentType = WindowsRuntime]
$null = [Windows.Media.Ocr.OcrEngine, Windows.Foundation, ContentType = WindowsRuntime]
$null = [Windows.Graphics.Imaging.BitmapDecoder, Windows.Graphics, ContentType = WindowsRuntime]
$asTask = ([System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object {
    $_.Name -eq 'AsTask' -and $_.GetParameters().Count -eq 1 -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncOperation`1' })[0]
function Await($op, $type) { $t = $asTask.MakeGenericMethod($type).Invoke($null, @($op)); $t.Wait(-1) | Out-Null; $t.Result }
$file = Await ([Windows.Storage.StorageFile]::GetFileFromPathAsync($Path)) ([Windows.Storage.StorageFile])
$stream = Await ($file.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream])
$dec = Await ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream)) ([Windows.Graphics.Imaging.BitmapDecoder])
$bmp = Await ($dec.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
$best = ''; $bestScore = -1
foreach ($lang in [Windows.Media.Ocr.OcrEngine]::AvailableRecognizerLanguages) {
    $eng = [Windows.Media.Ocr.OcrEngine]::TryCreateFromLanguage($lang)
    if ($eng -eq $null) { continue }
    $res = Await ($eng.RecognizeAsync($bmp)) ([Windows.Media.Ocr.OcrResult])
    $txt = ($res.Lines | ForEach-Object { $_.Text }) -join "`n"
    $score = ([regex]::Matches($txt, '\p{L}{2,}')).Count
    if ($score -gt $bestScore) { $best = $txt; $bestScore = $score }
}
Write-Output $best
"""


def image_ocr(src, p):
    """Copy text (OCR) with the text recognizer built into Windows 10/11, for every language installed."""
    import tempfile
    if not IS_WIN:
        raise ConversionError("Copying text from images needs Windows 10 or 11")
    Image, _ = _pil()
    img = _flatten(open_image(src))
    if max(img.size) > 9000:
        img.thumbnail((9000, 9000), Image.LANCZOS)
    tmp = Path(temp_dir())
    png = tmp / "ocr.png"
    img.save(str(png), format="PNG")
    script = tmp / "ocr.ps1"
    script.write_text(OCR_PS, encoding="utf-8-sig")
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                        "-File", str(script), "-Path", str(png)],
                       capture_output=True, timeout=120, stdin=subprocess.DEVNULL, creationflags=NO_WINDOW)
    text = r.stdout.decode("utf-8", "replace").strip()
    if r.returncode != 0 and not text:
        raise ConversionError("Windows text recognition isn't available on this PC")
    if not text:
        raise ConversionError("No text found in this image")
    out = out_for(src, "txt", "text")
    out.write_text(text + "\n", encoding="utf-8")
    return out


def images_to_pdf(paths, p):
    imgs = [_flatten(open_image(x)) for x in paths]
    out = out_for(paths[0], "pdf", "combined" if len(paths) > 1 else "")
    imgs[0].save(str(out), format="PDF", save_all=True, append_images=imgs[1:], resolution=150)
    return out


# ----------------------------------------------------------------------------
# PDF / EPUB  (PyMuPDF)
# ----------------------------------------------------------------------------
def _mu():
    import pymupdf
    return pymupdf


def pdf_to_images(fmt, dpi=200):
    def f(src, p):
        mu = _mu()
        doc = mu.open(str(src))
        if doc.page_count == 1:
            out = out_for(src, fmt)
            doc[0].get_pixmap(dpi=dpi).save(str(out), jpg_quality=92) if fmt == "jpg" else doc[0].get_pixmap(dpi=dpi).save(str(out))
            return out
        folder = out_dir_for(src, fmt.upper())
        folder.mkdir()
        width = len(str(doc.page_count))
        for i, page in enumerate(doc, 1):
            check_cancel()
            target = folder / f"{Path(src).stem} - {str(i).zfill(width)}.{fmt}"
            pix = page.get_pixmap(dpi=dpi)
            pix.save(str(target), jpg_quality=92) if fmt == "jpg" else pix.save(str(target))
        return folder
    return f


def pdf_to_text(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    out = out_for(src, "txt")
    out.write_text("\n\n".join(pg.get_text() for pg in doc), encoding="utf-8")
    return out


def pdf_to_markdown(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    parts = []
    for i, pg in enumerate(doc, 1):
        blocks = pg.get_text("dict")["blocks"]
        sizes = [s["size"] for b in blocks for l in b.get("lines", []) for s in l["spans"] if s["text"].strip()]
        body = sorted(sizes)[len(sizes) // 2] if sizes else 11
        for b in blocks:
            for l in b.get("lines", []):
                text = "".join(s["text"] for s in l["spans"]).strip()
                if not text:
                    continue
                size = max(s["size"] for s in l["spans"])
                bold = any(s["flags"] & 16 for s in l["spans"])
                if size >= body * 1.6:
                    parts.append(f"\n# {text}\n")
                elif size >= body * 1.25 or (bold and len(text) < 80 and size > body):
                    parts.append(f"\n## {text}\n")
                else:
                    parts.append(text)
            parts.append("")
    out = out_for(src, "md")
    out.write_text(re.sub(r"\n{3,}", "\n\n", "\n".join(parts)).strip() + "\n", encoding="utf-8")
    return out


def pdf_to_html(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    out = out_for(src, "html")
    body = "\n".join(pg.get_text("xhtml") for pg in doc)
    out.write_text(f"<!doctype html><meta charset='utf-8'><title>{Path(src).stem}</title><body>{body}</body>", encoding="utf-8")
    return out


def pdf_to_docx(src, p):
    try:
        from pdf2docx import Converter
    except ImportError:
        # fall back to Word, which can open PDFs
        return office_convert(src, "docx")
    out = out_for(src, "docx")
    cv = Converter(str(src))
    try:
        cv.convert(str(out))
    finally:
        cv.close()
    return out


def pdf_compress(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    try:
        doc.rewrite_images(dpi_threshold=160, dpi_target=150, quality=70, lossy=True, lossless=True, bitonal=True, color=True, gray=True)
    except Exception:
        pass
    try:
        doc.subset_fonts()
    except Exception:
        pass
    out = out_for(src, "pdf", "compressed")
    doc.save(str(out), garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
    return out


def pdf_split(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    folder = out_dir_for(src, "pages")
    folder.mkdir()
    width = len(str(doc.page_count))
    for i in range(doc.page_count):
        check_cancel()
        one = mu.open()
        one.insert_pdf(doc, from_page=i, to_page=i)
        one.save(str(folder / f"{Path(src).stem} - {str(i + 1).zfill(width)}.pdf"))
    return folder


def _page_list(spec: str, n: int) -> list[int]:
    pages = []
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            a, b = int(a or 1), int(b or n)
            pages += list(range(a, b + 1))
        else:
            pages.append(int(part))
    bad = [x for x in pages if not 1 <= x <= n]
    if bad or not pages:
        raise ConversionError(f"Pages must be between 1 and {n}")
    return [x - 1 for x in pages]


def pdf_extract_pages(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    doc.select(_page_list(p["Pages (e.g. 1-3,5)"], doc.page_count))
    out = out_for(src, "pdf", "pages")
    doc.save(str(out), garbage=3, deflate=True)
    return out


def pdf_rotate(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    for pg in doc:
        pg.set_rotation((pg.rotation + 90) % 360)
    out = out_for(src, "pdf", "rotated")
    doc.save(str(out))
    return out


def pdf_extract_images(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    folder = out_dir_for(src, "images")
    folder.mkdir()
    n = 0
    seen = set()
    for pi, pg in enumerate(doc, 1):
        for img in pg.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            info = doc.extract_image(xref)
            n += 1
            (folder / f"page{pi}-{n}.{info['ext']}").write_bytes(info["image"])
    if n == 0:
        folder.rmdir()
        raise ConversionError("No embedded images found")
    return folder


def pdf_merge(paths, p):
    mu = _mu()
    out_doc = mu.open()
    for x in paths:
        cat = EXT_CATEGORY.get(ext_of(x))
        if cat == "pdf":
            out_doc.insert_pdf(mu.open(str(x)))
        elif cat == "image":
            img = open_image(x)
            buf = io.BytesIO()
            _flatten(img).save(buf, format="PDF", resolution=150)
            out_doc.insert_pdf(mu.open("pdf", buf.getvalue()))
    out = out_for(paths[0], "pdf", "merged")
    out_doc.save(str(out), garbage=3, deflate=True)
    return out


def pdf_protect(src, p):
    pw = p["Password"]
    if not pw:
        raise ConversionError("Password cannot be empty")
    mu = _mu()
    doc = mu.open(str(src))
    out = out_for(src, "pdf", "protected")
    doc.save(str(out), encryption=mu.PDF_ENCRYPT_AES_256, user_pw=pw, owner_pw=pw)
    return out


def epub_to_pdf(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    out = out_for(src, "pdf")
    pdf = mu.open("pdf", doc.convert_to_pdf())
    pdf.save(str(out), garbage=3, deflate=True)
    return out


def epub_to_text(src, p):
    mu = _mu()
    doc = mu.open(str(src))
    out = out_for(src, "txt")
    out.write_text("\n\n".join(pg.get_text() for pg in doc), encoding="utf-8")
    return out


# ----------------------------------------------------------------------------
# Text documents (Markdown / TXT / HTML) — pure Python
# ----------------------------------------------------------------------------
HTML_CSS = """
body{font-family:'Segoe UI',Helvetica,Arial,sans-serif;font-size:11pt;line-height:1.5;color:#1d1d1f}
h1{font-size:22pt}h2{font-size:17pt}h3{font-size:13pt}
code,pre{font-family:Consolas,monospace;font-size:9.5pt;background:#f3f3f3}
pre{padding:6px}table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:3px 6px}
blockquote{color:#555;margin-left:12px}
"""


def _read_text(src) -> str:
    raw = Path(src).read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1256", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _to_html(src) -> str:
    e = ext_of(src)
    txt = _read_text(src)
    if e in ("md", "markdown"):
        import markdown
        return markdown.markdown(txt, extensions=["extra", "sane_lists", "tables"])
    if e in ("html", "htm"):
        m = re.search(r"<body[^>]*>(.*)</body>", txt, re.S | re.I)
        return m.group(1) if m else txt
    import html
    return "<pre style='white-space:pre-wrap;font-family:Consolas,monospace'>" + html.escape(txt) + "</pre>"


def html_to_pdf_file(html_body: str, out: Path) -> Path:
    mu = _mu()
    story = mu.Story(html=html_body, user_css=HTML_CSS)
    writer = mu.DocumentWriter(str(out))
    mediabox = mu.paper_rect("a4")
    where = mediabox + (54, 54, -54, -54)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return out


def text_to_pdf(src, p):
    return html_to_pdf_file(_to_html(src), out_for(src, "pdf"))


def text_to_html(src, p):
    out = out_for(src, "html")
    out.write_text(f"<!doctype html><meta charset='utf-8'><title>{Path(src).stem}</title>"
                   f"<style>{HTML_CSS}</style><body>{_to_html(src)}</body>", encoding="utf-8")
    return out


def html_to_markdown(src, p):
    from markdownify import markdownify
    out = out_for(src, "md")
    out.write_text(markdownify(_read_text(src), heading_style="ATX"), encoding="utf-8")
    return out


def md_to_txt(src, p):
    import html as h
    plain = re.sub(r"<[^>]+>", "", _to_html(src))
    out = out_for(src, "txt")
    out.write_text(h.unescape(plain).strip() + "\n", encoding="utf-8")
    return out


def text_to_docx(src, p):
    """MD/TXT/HTML -> DOCX: via Word/LibreOffice using an intermediate HTML file."""
    tmp = out_for(src, "html", "tmp")
    text_to_html_path = text_to_html(src, p)
    os.replace(text_to_html_path, tmp)
    try:
        res = office_convert(tmp, "docx")
        final = out_for(src, "docx")
        os.replace(res, final)
        return final
    finally:
        if tmp.exists():
            tmp.unlink()


# ----------------------------------------------------------------------------
# Office (Word / Excel / PowerPoint via COM, or LibreOffice)
# ----------------------------------------------------------------------------
def soffice() -> str | None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        return exe
    if IS_WIN:
        for base in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                     os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
            cand = Path(base) / "LibreOffice" / "program" / "soffice.exe"
            if cand.exists():
                return str(cand)
    return None


WORD_FMT = {"pdf": 17, "docx": 16, "rtf": 6, "txt": 7, "html": 10, "odt": 23, "doc": 0}
EXCEL_FMT = {"pdf": None, "xlsx": 51, "csv": 62, "ods": 60, "xls": 56}
PPT_FMT = {"pdf": 32, "pptx": 24, "png": 18, "jpg": 17, "odp": 35}


def _com_convert(src: Path, fmt: str) -> Path | None:
    if not IS_WIN:
        return None
    try:
        import pythoncom
        import win32com.client as wc
    except ImportError:
        return None
    cat = EXT_CATEGORY.get(ext_of(src))
    out = out_for(src, fmt)
    pythoncom.CoInitialize()
    try:
        if cat in ("document", "pdf") and fmt in WORD_FMT:
            app = wc.DispatchEx("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            try:
                doc = app.Documents.Open(str(src), ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False)
                doc.SaveAs2(str(out), FileFormat=WORD_FMT[fmt])
                doc.Close(False)
            finally:
                app.Quit()
            return out
        if cat == "spreadsheet" and fmt in EXCEL_FMT:
            app = wc.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            try:
                wb = app.Workbooks.Open(str(src), ReadOnly=True, UpdateLinks=0)
                if fmt == "pdf":
                    wb.ExportAsFixedFormat(0, str(out))
                else:
                    wb.SaveAs(str(out), FileFormat=EXCEL_FMT[fmt])
                wb.Close(False)
            finally:
                app.Quit()
            return out
        if cat == "slides" and fmt in PPT_FMT:
            app = wc.DispatchEx("PowerPoint.Application")
            try:
                pres = app.Presentations.Open(str(src), ReadOnly=True, WithWindow=False)
                if fmt in ("png", "jpg"):
                    out = out_dir_for(src, fmt.upper())
                pres.SaveAs(str(out), PPT_FMT[fmt])
                pres.Close()
            finally:
                app.Quit()
            return out
    except Exception as e:  # Office not installed or failed
        print("COM conversion failed:", e, file=sys.stderr)
        return None
    finally:
        pythoncom.CoUninitialize()
    return None


def _lo_convert(src: Path, fmt: str) -> Path | None:
    exe = soffice()
    if not exe:
        return None
    import tempfile
    filt = {"txt": "txt:Text (encoded):UTF8", "csv": "csv:Text - txt - csv (StarCalc):44,34,76",
            "docx": "docx:MS Word 2007 XML", "odt": "odt:writer8", "rtf": "rtf:Rich Text Format"}.get(fmt, fmt)
    cmd_extra = ["--infilter=HTML (StarWriter)"] if ext_of(src) in ("html", "htm") else []
    td = temp_dir()
    try:
        run([exe, "--headless", "--norestore", *cmd_extra, "--convert-to", filt, "--outdir", str(td), str(src)])
        produced = list(Path(td).glob(f"*.{fmt}"))
        if not produced:
            return None
        out = out_for(src, fmt)
        shutil.move(str(produced[0]), out)
        return out
    finally:
        shutil.rmtree(td, ignore_errors=True)


def office_available() -> bool:
    if soffice():
        return True
    if IS_WIN:
        try:
            import winreg
            for prog in ("Word.Application", "Excel.Application", "PowerPoint.Application"):
                try:
                    winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, prog).Close()
                    return True
                except OSError:
                    pass
        except Exception:
            pass
    return False


def office_convert(src, fmt) -> Path:
    src = Path(src).resolve()
    res = _com_convert(src, fmt) or _lo_convert(src, fmt)
    if not res:
        raise ConversionError("Needs Microsoft Office or LibreOffice installed for this conversion")
    return res


def office_to(fmt):
    return lambda src, p: office_convert(src, fmt)


def slides_to_images(fmt):
    def f(src, p):
        src = Path(src).resolve()
        if IS_WIN:
            res = _com_convert(src, fmt)
            if res:
                return res
        pdf = office_convert(src, "pdf")
        try:
            folder = pdf_to_images(fmt, dpi=150)(pdf, p)
            if folder.is_dir():
                final = out_dir_for(src, fmt.upper())
                folder.rename(final)
                return final
            final = out_for(src, fmt)
            folder.rename(final)
            return final
        finally:
            pdf.unlink(missing_ok=True)
    return f


# ----------------------------------------------------------------------------
# Spreadsheets (openpyxl / csv)
# ----------------------------------------------------------------------------
def _read_rows(src) -> list[tuple[str, list[list]]]:
    e = ext_of(src)
    if e in ("csv", "tsv"):
        txt = _read_text(src)
        delim = "\t" if e == "tsv" else (csv.Sniffer().sniff(txt[:4096], ",;\t|").delimiter if txt.strip() else ",")
        return [(Path(src).stem, list(csv.reader(io.StringIO(txt), delimiter=delim)))]
    if e in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(str(src), read_only=True, data_only=True)
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    # xls / ods -> convert via office first
    tmp = office_convert(src, "xlsx")
    try:
        return _read_rows(tmp)
    finally:
        tmp.unlink(missing_ok=True)


def _clean(v):
    return "" if v is None else v


def sheet_to_delimited(fmt):
    def f(src, p):
        sheets = _read_rows(src)
        delim = "\t" if fmt == "tsv" else ","
        if len(sheets) == 1:
            out = out_for(src, fmt)
            with open(out, "w", newline="", encoding="utf-8-sig") as fh:
                csv.writer(fh, delimiter=delim).writerows([[_clean(v) for v in r] for r in sheets[0][1]])
            return out
        folder = out_dir_for(src, fmt.upper())
        folder.mkdir()
        for title, rows in sheets:
            safe = re.sub(r'[\\/:*?"<>|]', "_", title)
            with open(folder / f"{safe}.{fmt}", "w", newline="", encoding="utf-8-sig") as fh:
                csv.writer(fh, delimiter=delim).writerows([[_clean(v) for v in r] for r in rows])
        return folder
    return f


def sheet_to_xlsx(src, p):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in _read_rows(src):
        ws = wb.create_sheet(title[:31] or "Sheet1")
        for r in rows:
            ws.append([_num(v) for v in r])
    out = out_for(src, "xlsx")
    wb.save(str(out))
    return out


def _num(v):
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r"-?\d+", s) and len(s) < 15 and not (len(s) > 1 and s.startswith("0")):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            try:
                return float(s)
            except ValueError:
                pass
    return v


def sheet_to_json(src, p):
    sheets = _read_rows(src)
    data = {}
    for title, rows in sheets:
        if not rows:
            data[title] = []
            continue
        head = [str(h) if h is not None else f"col{i + 1}" for i, h in enumerate(rows[0])]
        data[title] = [{head[i]: (_num(v) if isinstance(v, str) else v) for i, v in enumerate(r) if i < len(head)} for r in rows[1:]]
    out = out_for(src, "json")
    payload = next(iter(data.values())) if len(data) == 1 else data
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return out


def sheet_to_html(src, p):
    import html as h
    parts = [f"<!doctype html><meta charset='utf-8'><title>{h.escape(Path(src).stem)}</title><style>{HTML_CSS}</style><body>"]
    for title, rows in _read_rows(src):
        parts.append(f"<h2>{h.escape(title)}</h2><table>")
        for i, r in enumerate(rows):
            tag = "th" if i == 0 else "td"
            parts.append("<tr>" + "".join(f"<{tag}>{h.escape(str(_clean(v)))}</{tag}>" for v in r) + "</tr>")
        parts.append("</table>")
    out = out_for(src, "html")
    out.write_text("".join(parts), encoding="utf-8")
    return out


def sheet_to_markdown(src, p):
    def cell(v):
        return str(_clean(v)).replace("|", "\\|").replace("\n", " ")
    parts = []
    sheets = _read_rows(src)
    for title, rows in sheets:
        rows = [r for r in rows if any(v not in (None, "") for v in r)]
        if len(sheets) > 1:
            parts.append(f"## {title}\n")
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [list(r) + [""] * (width - len(r)) for r in rows]
        parts.append("| " + " | ".join(cell(v) for v in rows[0]) + " |")
        parts.append("|" + "---|" * width)
        for r in rows[1:]:
            parts.append("| " + " | ".join(cell(v) for v in r) + " |")
        parts.append("")
    out = out_for(src, "md")
    out.write_text("\n".join(parts), encoding="utf-8")
    return out


def sheet_to_pdf(src, p):
    if office_available():
        try:
            return office_convert(src, "pdf")
        except ConversionError:
            pass
    tmp = sheet_to_html(src, p)
    try:
        return html_to_pdf_file(_to_html(tmp), out_for(src, "pdf"))
    finally:
        tmp.unlink(missing_ok=True)


# ----------------------------------------------------------------------------
# Subtitles
# ----------------------------------------------------------------------------
def srt_to_vtt(src, p):
    txt = _read_text(src).replace("\r\n", "\n")
    txt = re.sub(r"(\d{2}:\d{2}:\d{2}),(\d{3})", r"\1.\2", txt)
    out = out_for(src, "vtt")
    out.write_text("WEBVTT\n\n" + txt.strip() + "\n", encoding="utf-8")
    return out


def vtt_to_srt(src, p):
    txt = _read_text(src).replace("\r\n", "\n")
    txt = re.sub(r"^WEBVTT.*?\n\n", "", txt, flags=re.S)
    cues = [c for c in re.split(r"\n{2,}", txt.strip()) if "-->" in c]
    out_lines = []
    for i, cue in enumerate(cues, 1):
        lines = cue.split("\n")
        while lines and "-->" not in lines[0]:
            lines.pop(0)
        ts = lines[0]
        ts = re.sub(r"(?<![\d:])(\d{2}:\d{2})\.(\d{3})", r"00:\1,\2", ts)
        ts = re.sub(r"(\d{2}:\d{2}:\d{2})\.(\d{3})", r"\1,\2", ts)
        ts = re.sub(r"\s+(align|line|position|size|vertical):\S+", "", ts)
        out_lines.append(f"{i}\n{ts}\n" + "\n".join(lines[1:]))
    out = out_for(src, "srt")
    out.write_text("\n\n".join(out_lines) + "\n", encoding="utf-8")
    return out


def subs_ffmpeg(fmt):
    def f(src, p):
        return ff(src, out_for(src, fmt))
    return f


def subs_to_txt(src, p):
    txt = _read_text(src).replace("\r\n", "\n")
    keep = []
    if ext_of(src) in ("ass", "ssa"):
        for line in txt.split("\n"):
            if line.startswith("Dialogue:"):
                t = line.split(",", 9)[-1]
                t = re.sub(r"\{[^}]*\}", "", t).replace("\\N", "\n").replace("\\n", "\n").strip()
                if t and (not keep or keep[-1] != t):
                    keep.append(t)
        out = out_for(src, "txt")
        out.write_text("\n".join(keep) + "\n", encoding="utf-8")
        return out
    for line in txt.split("\n"):
        s = line.strip()
        if not s or s == "WEBVTT" or s.isdigit() or "-->" in s:
            continue
        s = re.sub(r"<[^>]+>", "", s)
        if not keep or keep[-1] != s:
            keep.append(s)
    out = out_for(src, "txt")
    out.write_text("\n".join(keep) + "\n", encoding="utf-8")
    return out


# ----------------------------------------------------------------------------
# Archives
#   read : zip, 7z, tar, tar.gz/tgz, tar.bz2/tbz2, tar.xz/txz, gz, bz2, xz  (built in)
#          rar (+ iso, cab…) via 7-Zip, WinRAR, or Windows' own tar.exe
#   write: zip, 7z, tar, tar.gz, tar.bz2, tar.xz (built in), rar (needs WinRAR)
# ----------------------------------------------------------------------------
ARCHIVE_OUT = ["zip", "7z", "rar", "tar", "tar.gz", "tar.bz2", "tar.xz"]
ARCHIVE_LABEL = {"zip": "ZIP", "7z": "7Z", "rar": "RAR", "tar": "TAR", "tar.gz": "TAR.GZ",
                 "tar.bz2": "TAR.BZ2", "tar.xz": "TAR.XZ"}
# ext_of() value -> canonical archive format
ARCHIVE_FMT = {"zip": "zip", "7z": "7z", "rar": "rar", "tar": "tar", "tgz": "tar.gz", "tbz2": "tar.bz2",
               "txz": "tar.xz", "gz": "gz", "bz2": "bz2", "xz": "xz", "iso": "iso", "cab": "cab"}


def _program_files():
    return [os.environ.get("ProgramFiles", r"C:\Program Files"),
            os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
            os.environ.get("ProgramW6432", r"C:\Program Files")]


def seven_zip() -> str | None:
    exe = shutil.which("7z") or shutil.which("7za")
    if exe:
        return exe
    if IS_WIN:
        for base in _program_files():
            cand = Path(base) / "7-Zip" / "7z.exe"
            if cand.exists():
                return str(cand)
    return None


def winrar(tool="Rar.exe") -> str | None:
    """Rar.exe (create) or UnRAR.exe (extract) from a WinRAR install."""
    exe = shutil.which(tool.split(".")[0].lower())
    if exe:
        return exe
    if IS_WIN:
        for base in _program_files():
            cand = Path(base) / "WinRAR" / tool
            if cand.exists():
                return str(cand)
    return None


def windows_tar() -> str | None:
    """Windows 10 1803+ ships bsdtar (libarchive) which can read rar/7z/iso/cab."""
    if IS_WIN:
        cand = Path(os.environ.get("SystemRoot", r"C:\Windows")) / "System32" / "tar.exe"
        if cand.exists():
            return str(cand)
    return None


def can_write_rar() -> bool:
    return winrar("Rar.exe") is not None


def archive_stem(path: Path) -> str:
    n = path.name
    for suf in (".tar.gz", ".tar.bz2", ".tar.xz"):
        if n.lower().endswith(suf):
            return n[: -len(suf)]
    return path.stem if path.is_file() or path.suffix else path.name


def _safe(dest: Path, name: str):
    target = (dest / name).resolve()
    if not str(target).startswith(str(dest.resolve())):
        raise ConversionError(f"Unsafe path in archive: {name}")


def _extract_into(src: Path, dest: Path):
    fmt = ARCHIVE_FMT.get(ext_of(src), "")
    if fmt == "zip":
        with zipfile.ZipFile(src) as z:
            for m in z.namelist():
                _safe(dest, m)
            for m in z.infolist():
                check_cancel()
                z.extract(m, dest)
        return
    if fmt == "7z":
        import py7zr
        with py7zr.SevenZipFile(src, "r") as z:
            for m in z.getnames():
                _safe(dest, m)
            z.extractall(dest)
        return
    if fmt.startswith("tar") or (fmt in ("gz", "bz2", "xz") and tarfile.is_tarfile(src)):
        with tarfile.open(src) as t:
            for m in t.getmembers():
                _safe(dest, m.name)
            t.extractall(dest, filter="data") if hasattr(tarfile, "data_filter") else t.extractall(dest)
        return
    if fmt in ("gz", "bz2", "xz"):
        import bz2
        import gzip
        import lzma
        opener = {"gz": gzip.open, "bz2": bz2.open, "xz": lzma.open}[fmt]
        with opener(src, "rb") as fi, open(dest / src.stem, "wb") as fo:
            shutil.copyfileobj(fi, fo)
        return
    # rar, iso, cab, anything else: external tools, best first
    if (exe := seven_zip()):
        run([exe, "x", "-y", f"-o{dest}", str(src)])
    elif fmt == "rar" and (exe := winrar("UnRAR.exe")):
        run([exe, "x", "-y", "-idq", str(src), str(dest) + os.sep])
    elif (exe := windows_tar()):
        run([exe, "-xf", str(src), "-C", str(dest)])
    else:
        raise ConversionError("Install 7-Zip or WinRAR to open this archive")


def archive_extract(src, p):
    src = Path(src)
    dest = unique_path(_target_dir(src) / archive_stem(src))
    dest.mkdir()
    try:
        _extract_into(src, dest)
    except Exception:
        shutil.rmtree(dest, ignore_errors=True)
        raise
    # collapse a single top-level folder
    items = list(dest.iterdir())
    if len(items) == 1 and items[0].is_dir():
        inner = items[0]
        tmp = dest.with_name(dest.name + ".__tmp")
        inner.rename(tmp)
        dest.rmdir()
        final = unique_path(dest.parent / inner.name)
        tmp.rename(final)
        return final
    return dest


def archive_repack(fmt):
    """Convert one archive format to another (extract to temp, repack)."""
    def f(src, p):
        src = Path(src)
        td = temp_dir()
        try:
            _extract_into(src, td)
            out = unique_path(_target_dir(src) / f"{archive_stem(src)}.{fmt}")
            return _write_archive(sorted(td.iterdir()), td, fmt, out)
        finally:
            shutil.rmtree(td, ignore_errors=True)
    return f


def _write_archive(items: list[Path], root: Path, fmt: str, out: Path) -> Path:
    """Write items (files/dirs) into out, storing paths relative to root."""
    out = Path(os.path.abspath(out))
    root = Path(os.path.abspath(root))
    items = [Path(os.path.abspath(it)) for it in items]
    tmp = out.with_name(out.name + ".part")
    try:
        if fmt == "zip":
            with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
                for it in items:
                    if it.is_dir():
                        z.write(it, it.relative_to(root))
                        for f in sorted(it.rglob("*")):
                            check_cancel()
                            z.write(f, f.relative_to(root))
                    else:
                        z.write(it, it.relative_to(root))
        elif fmt == "7z":
            import py7zr
            with py7zr.SevenZipFile(tmp, "w") as z:
                for it in items:
                    if it.is_dir():
                        for f in [it, *sorted(it.rglob("*"))]:
                            check_cancel()
                            z.write(f, str(f.relative_to(root)))
                    else:
                        check_cancel()
                        z.write(it, str(it.relative_to(root)))
        elif fmt.startswith("tar"):
            mode = {"tar": "w", "tar.gz": "w:gz", "tar.bz2": "w:bz2", "tar.xz": "w:xz"}[fmt]
            with tarfile.open(tmp, mode) as t:
                for it in items:
                    check_cancel()
                    t.add(it, arcname=str(it.relative_to(root)),
                          filter=lambda ti: (check_cancel(), ti)[1])
        elif fmt == "rar":
            exe = winrar("Rar.exe")
            if not exe:
                raise ConversionError("Creating RAR files needs WinRAR installed")
            rel = [str(it.relative_to(root)) for it in items]
            run([exe, "a", "-r", "-idq", "-y", str(tmp), *rel], cwd=str(root), ok_codes=(0, 1))
            # Rar.exe appends .rar when the name doesn't end with it
            if not tmp.exists() and Path(str(tmp) + ".rar").exists():
                Path(str(tmp) + ".rar").rename(tmp)
        else:
            raise ConversionError(f"Unsupported archive format: {fmt}")
        os.replace(tmp, out)
        return out
    except Exception:
        tmp.unlink(missing_ok=True)
        raise


def pack_paths(paths, fmt) -> Path:
    """Compress files/folders into one archive next to the first item."""
    items = [Path(x) for x in paths]
    first = items[0]
    if len(items) == 1:
        name = first.name if first.is_dir() else first.stem
    else:
        name = first.parent.name or "Archive"
    out = unique_path(_target_dir(first) / f"{name}.{fmt}")
    return _write_archive(items, first.parent, fmt, out)


def compress_to(fmt):
    return lambda paths, p: pack_paths(paths, fmt)


def pack(folder_or_files, fmt: str, name_like: Path | None = None) -> Path:   # kept for older callers
    items = folder_or_files if isinstance(folder_or_files, list) else [folder_or_files]
    return pack_paths(items, fmt)


def to_zip(paths, p):
    return pack_paths(paths, "zip")


def to_targz(paths, p):
    return pack_paths(paths, "tar.gz")


def gzip_file(src, p):
    import gzip
    out = unique_path(Path(str(src) + ".gz"))
    with open(src, "rb") as fi, gzip.open(out, "wb", 9) as fo:
        shutil.copyfileobj(fi, fo)
    return out


def archive_targets(exclude: str = "") -> list[str]:
    return [f for f in ARCHIVE_OUT if f != exclude]


def _rar_hint(f):
    return "needs WinRAR" if f == "rar" and not can_write_rar() else ""


# ----------------------------------------------------------------------------
# Catalog
# ----------------------------------------------------------------------------
@dataclass
class Op:
    id: str
    label: str
    fn: Callable
    group: str = "convert"           # convert | tool | action
    ask: list[tuple[str, str]] = field(default_factory=list)  # [(field, default)]
    multi: bool = False              # fn receives the whole list of paths
    hint: str = ""
    ui: str = ""                     # "crop" -> the app opens the crop & resize editor first
    model: str = ""                  # AI model the op needs (downloaded once, with consent)


IMAGE_OUT = ["png", "jpg", "webp", "heic", "avif", "tiff", "bmp", "gif", "ico", "tga", "jp2", "icns", "pdf"]
VIDEO_OUT = ["mp4", "mov", "mkv", "webm", "avi", "wmv"]
AUDIO_OUT = ["mp3", "m4a", "wav", "flac", "ogg", "opus", "aiff", "wma", "ac3"]
SUB_OUT = ["srt", "vtt", "ass"]


def _norm(e):
    return {"jpeg": "jpg", "tif": "tiff", "heif": "heic", "htm": "html", "markdown": "md",
            "aif": "aiff", "oga": "ogg", "m4v": "mp4", "j2k": "jp2", "jpx": "jp2", "ssa": "ass"}.get(e, e)


def conversions(paths: list[str]) -> list[Op]:
    cat = category_of(paths)
    e = _norm(ext_of(paths[0])) if paths else ""
    ops: list[Op] = []
    add = lambda fmt, fn, label=None: ops.append(Op(f"to_{fmt}", label or fmt.upper(), fn)) if fmt != e else None

    if cat == "image":
        for f in IMAGE_OUT:
            if f == "pdf":
                ops.append(Op("to_pdf", "PDF", images_to_pdf, multi=True))
            else:
                add(f, image_to(f))
    elif cat == "video":
        for f in VIDEO_OUT:
            add(f, video_to(f))
        ops.append(Op("to_hevc", "MP4 H.265", video_to_hevc, hint="smaller file"))
        add("gif", video_to_gif)
        ops.append(Op("to_mp3", "MP3", audio_to("mp3"), hint="audio only"))
        ops.append(Op("to_m4a", "M4A", audio_to("m4a"), hint="audio only"))
    elif cat == "audio":
        for f in AUDIO_OUT:
            add(f, audio_to(f))
    elif cat == "pdf":
        add("docx", pdf_to_docx, "DOCX")
        add("png", pdf_to_images("png"))
        add("jpg", pdf_to_images("jpg"))
        add("txt", pdf_to_text)
        add("md", pdf_to_markdown, "Markdown")
        add("html", pdf_to_html)
    elif cat == "document":
        if e in ("md", "txt", "html"):
            add("pdf", text_to_pdf)
            add("html", text_to_html)
            if e == "html":
                add("md", html_to_markdown, "Markdown")
            if e == "md":
                add("txt", md_to_txt)
            if office_available():
                add("docx", text_to_docx)
        else:  # docx / doc / rtf / odt
            for f in ("pdf", "docx", "odt", "rtf", "txt", "html"):
                add(f, office_to(f))
    elif cat == "slides":
        add("pdf", office_to("pdf"))
        add("png", slides_to_images("png"))
        add("jpg", slides_to_images("jpg"))
        if e != "pptx":
            add("pptx", office_to("pptx"))
        add("odp", office_to("odp"))
    elif cat == "spreadsheet":
        add("xlsx", sheet_to_xlsx)
        add("csv", sheet_to_delimited("csv"))
        add("tsv", sheet_to_delimited("tsv"))
        add("json", sheet_to_json)
        add("html", sheet_to_html)
        add("md", sheet_to_markdown, "Markdown")
        add("pdf", sheet_to_pdf)
        if office_available():
            add("ods", office_to("ods"))
    elif cat == "subtitle":
        for f in SUB_OUT:
            if f == e:
                continue
            if e == "srt" and f == "vtt":
                add("vtt", srt_to_vtt)
            elif e == "vtt" and f == "srt":
                add("srt", vtt_to_srt)
            else:
                add(f, subs_ffmpeg(f))
        add("txt", subs_to_txt)
    elif cat == "ebook":
        add("pdf", epub_to_pdf)
        add("txt", epub_to_text)
        add("md", pdf_to_markdown, "Markdown")
        add("html", pdf_to_html)
        if e in ("cbz", "xps", "oxps"):
            add("png", pdf_to_images("png", dpi=150))
            add("jpg", pdf_to_images("jpg", dpi=150))
    elif cat == "archive":                 # "Extract" lives in Tools
        cur = ARCHIVE_FMT.get(ext_of(paths[0]), "")
        for f in archive_targets(exclude=cur):
            ops.append(Op(f"to_{f}", ARCHIVE_LABEL[f], archive_repack(f), hint=_rar_hint(f)))
    elif cat in ("folder", "mixed", "other"):
        for f in archive_targets():
            ops.append(Op(f"to_{f}", ARCHIVE_LABEL[f], compress_to(f), multi=True, hint=_rar_hint(f)))
        if cat == "other" and len(paths) == 1:
            ops.append(Op("to_gz", "GZIP", gzip_file))
    return ops


TRIM = [("Start", "0:00"), ("End", "0:30")]


def _user_name() -> str:
    try:
        import getpass
        return getpass.getuser()
    except Exception:
        return ""


def tools(paths: list[str]) -> list[Op]:
    cat = category_of(paths)
    n = len(paths)
    T = lambda *a, **k: Op(*a, group="tool", **k)
    ops: list[Op] = []
    if cat == "image":
        ops += [
            T("compress", "Compress", image_compress),
            T("edit", "Crop & Resize…", image_crop_resize, ui="crop"),
            T("rotate", "Rotate 90°", image_tool("rotated", lambda i, p: i.rotate(-90, expand=True))),
            T("flip", "Flip", image_tool("flipped", lambda i, p: i.transpose(0))),
            T("gray", "Grayscale", image_tool("gray", lambda i, p: i.convert("LA") if i.mode in ("RGBA", "LA") else i.convert("L"))),
            T("strip", "Strip metadata", image_tool("clean", _strip)),
            T("resize", "Resize…", image_resize_preset, ask=[("Output size", RESIZE_PRESETS)]),
            T("removebg", "Remove background", image_remove_bg, model="isnet"),
            T("ocr", "Copy text", image_ocr),
            T("upscale", "Upscale 2×", image_upscale),
            T("sharpen", "Sharpen", image_sharpen),
            T("watermark", "Watermark…", image_watermark, ask=[("Text", "© " + _user_name()), ("Position", WATERMARK_POSITIONS)]),
            T("blurfaces", "Blur faces", image_blur_faces),
        ]
        if n > 1:
            ops.append(T("combine_pdf", "Combine PDF", images_to_pdf, multi=True))
            ops.append(T("stitch_h", "Stitch side by side", images_stitch("h"), multi=True))
            ops.append(T("stitch_v", "Stitch top to bottom", images_stitch("v"), multi=True))
            ops.append(T("gif", "Make GIF…", images_to_gif, multi=True, ask=[("Seconds per image", "0.5")]))
    elif cat == "video":
        ops += [
            T("compress", "Compress", video_compress),
            T("720p", "720p", video_scale(720)),
            T("trim", "Trim…", media_trim, ask=TRIM),
            T("mute", "Mute", video_mute),
            T("rotate", "Rotate 90°", video_rotate),
            T("snapshot", "Snapshot…", video_snapshot, ask=[("At time", "0:01")]),
            T("speed", "Speed 2×", video_speed2x),
            T("gif", "Make GIF…", video_gif_clip, ask=[("Start", "0:00"), ("End", "0:05")]),
        ]
        if n > 1:
            ops.append(T("merge", "Merge", media_merge, multi=True))
    elif cat == "audio":
        ops += [
            T("compress", "Compress", audio_compress),
            T("trim", "Trim…", media_trim, ask=TRIM),
            T("normalize", "Normalize", audio_normalize),
            T("mono", "Mono", audio_mono),
            T("fade", "Fade in/out", audio_fade),
        ]
        if n > 1:
            ops.append(T("merge", "Merge", media_merge, multi=True))
    elif cat == "pdf" or (cat == "mixed" and all(EXT_CATEGORY.get(ext_of(x)) in ("pdf", "image") for x in paths)):
        if cat == "pdf":
            ops += [
                T("compress", "Compress", pdf_compress),
                T("split", "Split pages", pdf_split),
                T("pages", "Extract pages…", pdf_extract_pages, ask=[("Pages (e.g. 1-3,5)", "1")]),
                T("rotate", "Rotate 90°", pdf_rotate),
                T("images", "Extract images", pdf_extract_images),
                T("protect", "Password…", pdf_protect, ask=[("Password", "")]),
            ]
        if n > 1:
            ops.append(T("merge", "Merge", pdf_merge, multi=True))
    elif cat == "archive":
        ops.append(T("extract", "Extract", archive_extract))
    # Folders, archives and "other" files already get ZIP/7Z/... in the Convert wheel
    if cat not in ("archive", "folder", "mixed", "other"):
        ops.append(T("zip", "Zip", to_zip, multi=True))
    return ops


def run_op(op: Op, paths: list[str], params: dict | None = None,
           token: CancelToken | None = None) -> list[Path]:
    params = params or {}
    _ctx.token = token
    try:
        if op.multi:
            outs = [op.fn(paths, params)]
        else:
            outs = []
            for p in paths:
                check_cancel()
                outs.append(op.fn(p, params))
        check_cancel()            # cancelled while a non-interruptible step was finishing
        return outs
    except BaseException:
        if token is not None and token.cancelled:
            _cleanup(token)
            raise Cancelled()
        raise
    finally:
        _ctx.token = None


def _cleanup(token: CancelToken):
    for p in reversed(token.paths):
        try:
            for cand in (p, p.with_name(p.name + ".part")):
                if cand.is_dir():
                    shutil.rmtree(cand, ignore_errors=True)
                elif cand.exists():
                    cand.unlink()
        except Exception:
            pass


def count_catalog() -> tuple[int, int]:
    """Rough count of conversion pairs and tools, for the About box."""
    samples = ["a.png", "a.mp4", "a.mp3", "a.pdf", "a.md", "a.docx", "a.html", "a.txt", "a.pptx",
               "a.xlsx", "a.csv", "a.srt", "a.vtt", "a.epub", "a.zip", "a.tar"]
    conv = sum(len(conversions([s])) for s in samples)
    tl = sum(len(tools([s])) for s in samples)
    return conv, tl
