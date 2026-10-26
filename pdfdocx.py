"""PDF → Word that looks the same ("exact layout").

Every line of text becomes a small positioned text frame at the same place, in the same font, size, weight and
colour; tables become real Word tables at the same place; lines, boxes, logos and pictures become one picture
behind the text. Arabic and other right-to-left text is written as proper right-to-left paragraphs, so it can be
edited in Word like any other document.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

RTL_RE = re.compile(r"[֐-ࣿיִ-﷿ﹰ-﻿]")
TW = 20          # twips per point

_FONT_NAMES = {
    "timesnewroman": "Times New Roman", "timesnewromanps": "Times New Roman", "arial": "Arial",
    "arialnarrow": "Arial Narrow", "arialblack": "Arial Black", "couriernew": "Courier New",
    "simplifiedarabic": "Simplified Arabic", "traditionalarabic": "Traditional Arabic",
    "arabictypesetting": "Arabic Typesetting", "sakkalmajalla": "Sakkal Majalla", "segoeui": "Segoe UI",
    "andalus": "Andalus", "tahoma": "Tahoma", "calibri": "Calibri", "calibrilight": "Calibri Light",
    "cambria": "Cambria", "aldhabi": "Aldhabi", "almohanad": "AL-Mohanad", "almohanadbold": "AL-Mohanad Bold",
    "diwanikufi": "Diwani Kufi", "microsoftuighur": "Microsoft Uighur", "urdtypesetting": "Urdu Typesetting",
    "garamond": "Garamond", "georgia": "Georgia", "verdana": "Verdana", "trebuchetms": "Trebuchet MS",
    "symbol": "Symbol", "wingdings": "Wingdings", "carlito": "Calibri", "liberationsans": "Arial",
    "liberationserif": "Times New Roman", "liberationmono": "Courier New", "dejavusans": "DejaVu Sans",
    "dejavuserif": "DejaVu Serif",
}


def _family(font: str) -> tuple[str, bool, bool]:
    """'ABCDEF+SimplifiedArabic-Bold' → ('Simplified Arabic', bold, italic)."""
    name = font.split("+", 1)[-1]
    low = name.lower()
    bold = any(k in low for k in ("bold", "black", "heavy", "semibold", "demi"))
    italic = "italic" in low or "oblique" in low
    base = re.split(r"[-,]", name)[0]
    base = re.sub(r"(PSMT|MT|PS)$", "", base)
    key = re.sub(r"[^a-z]", "", base.lower())
    for suffix in ("bold", "italic", "regular"):
        if key.endswith(suffix) and key[: -len(suffix)] in _FONT_NAMES:
            key = key[: -len(suffix)]
    return _FONT_NAMES.get(key, base), bold, italic


def _fix_ar(text: str) -> str:
    try:
        import engine
        return engine._fix_arabic_ligatures(text)
    except Exception:
        return text


# ----------------------------------------------------------------------------------------------- reading the PDF
def _spans(page):
    import pymupdf
    out = []
    lid = 0
    d = page.get_text("dict", flags=pymupdf.TEXT_PRESERVE_WHITESPACE | pymupdf.TEXT_PRESERVE_LIGATURES)
    for b in d["blocks"]:
        if b.get("type") != 0:
            continue
        for ln in b["lines"]:
            if abs(ln["dir"][1]) > 0.2:            # rotated text stays in the picture
                continue
            lid += 1
            for seq, s in enumerate(ln["spans"]):
                if not s["text"] or s["size"] < 2:
                    continue
                fam, bold, italic = _family(s["font"])
                out.append({
                    "text": s["text"], "bbox": tuple(s["bbox"]), "size": s["size"], "font": fam,
                    "bold": bold or bool(s["flags"] & 16), "italic": italic or bool(s["flags"] & 2),
                    "color": s["color"], "base": s["origin"][1], "lid": lid, "seq": seq,
                })
    return out


def _good_tables(page, spans):
    """Real tables only: ruled grids with text in them (not the frame lines around a letter)."""
    found = []
    try:
        tabs = page.find_tables().tables
    except Exception:
        return found
    for t in tabs:
        rows = [r.cells for r in t.rows]
        if len(rows) < 2 or max(len(r) for r in rows) < 2:
            continue
        cells = [c for r in rows for c in r]
        real = [c for c in cells if c]
        if len(real) < 0.7 * len(cells):
            continue
        if max(c[3] - c[1] for c in real) > 140:
            continue
        with_text = sum(1 for c in real if any(_inside(s["bbox"], c) for s in spans if s["text"].strip()))
        if with_text < 0.4 * len(real):
            continue
        found.append((tuple(t.bbox), rows))
    return found


def _inside(bb, rect, slack=1.5):
    cx, cy = (bb[0] + bb[2]) / 2, (bb[1] + bb[3]) / 2
    return rect[0] - slack <= cx <= rect[2] + slack and rect[1] - slack <= cy <= rect[3] + slack


_NEUTRAL = ".:،,؛;!?"


def _neutral_to_end(runs):
    """Some PDFs keep the full stop or colon of an Arabic line at the start of the leftmost piece
    (".هواتفهم"): put it back after the word."""
    last = runs[-1]
    t = last["text"]
    core = t.lstrip()
    if core and core[0] in _NEUTRAL and len(core) > 1 and RTL_RE.match(core[1:].lstrip() or " "):
        lead = t[: len(t) - len(core)]
        runs = runs[:-1] + [dict(last, text=lead + core[1:] + core[0])]
    return runs


def _pieces(spans):
    """The PDF's own lines (their pieces keep the order the PDF stores them in, which is the reading order),
    split where a line jumps across the page."""
    by_line: dict[int, list[dict]] = {}
    for s in spans:
        by_line.setdefault(s["lid"], []).append(s)
    pieces = []
    for group in by_line.values():
        group.sort(key=lambda s: s["seq"])
        cur: list[dict] = []
        last = None                                   # the last span with something visible
        for s in group:
            if s["text"].strip() and last is not None:
                dist = max(s["bbox"][0] - last["bbox"][2], last["bbox"][0] - s["bbox"][2])
                if dist > 2.2 * max(s["size"], last["size"]):
                    pieces.append(cur)
                    cur = []
            cur.append(s)
            if s["text"].strip():
                last = s
        pieces.append(cur)
    out = []
    for runs in pieces:
        vis = [s for s in runs if s["text"].strip()]
        if not vis:
            continue
        while not runs[0]["text"].strip():
            runs = runs[1:]
        while not runs[-1]["text"].strip():
            runs = runs[:-1]
        out.append({"runs": runs, "size": max(s["size"] for s in vis), "base": vis[0]["base"],
                    "bbox": (min(s["bbox"][0] for s in vis), min(s["bbox"][1] for s in vis),
                             max(s["bbox"][2] for s in vis), max(s["bbox"][3] for s in vis))})
    return out


def _lines(spans):
    """Text on the same baseline that belongs together, in reading order (right to left for Arabic)."""
    pieces = sorted(_pieces(spans), key=lambda p: (round(p["base"], 1), p["bbox"][0]))
    rows: list[list[dict]] = []
    for p in pieces:
        for r in rows:
            ref = r[0]
            if abs(ref["base"] - p["base"]) < 0.35 * max(ref["size"], p["size"]):
                r.append(p)
                break
        else:
            rows.append([p])
    segs = []
    for r in rows:
        r.sort(key=lambda p: p["bbox"][0])
        seg = [r[0]]
        for p in r[1:]:
            gap = p["bbox"][0] - max(q["bbox"][2] for q in seg)
            if gap > 2.2 * max(p["size"], seg[-1]["size"]):
                segs.append(seg)
                seg = [p]
            else:
                seg.append(p)
        segs.append(seg)
    lines = []
    for seg in segs:
        text = "".join(s["text"] for p in seg for s in p["runs"])
        rtl = bool(RTL_RE.search(text))
        ordered = sorted(seg, key=lambda p: -(p["bbox"][0] + p["bbox"][2])) if rtl else seg
        runs = []
        prev = None
        for p in ordered:
            if prev is not None:
                gap = (prev["bbox"][0] - p["bbox"][2]) if rtl else (p["bbox"][0] - prev["bbox"][2])
                if gap > 0.18 * p["size"] and not runs[-1]["text"].endswith(" ") and not p["runs"][0]["text"].startswith(" "):
                    runs.append(dict(runs[-1], text=" "))
            runs.extend(p["runs"])
            prev = p
        if rtl:
            runs = _neutral_to_end(runs)
        lines.append({"runs": runs, "rtl": rtl, "size": max(p["size"] for p in seg),
                      "bbox": (min(p["bbox"][0] for p in seg), min(p["bbox"][1] for p in seg),
                               max(p["bbox"][2] for p in seg), max(p["bbox"][3] for p in seg))})
    lines.sort(key=lambda ln: (ln["bbox"][1], ln["bbox"][0]))
    return lines


# ----------------------------------------------------------------------------------------------- writing Word
def _el(tag, **attrs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":", 1)), str(v))
    return e


_SEQ = {
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
           "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl "
           "cs em lang eastAsianLayout specVanish oMath",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
           "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
}


def _order(el):
    """Word refuses files whose properties are out of the schema's order: put them back in order."""
    seq = _SEQ.get(el.tag.rsplit("}", 1)[-1])
    if not seq:
        return
    rank = {n: i for i, n in enumerate(seq.split())}
    kids = list(el)
    kids.sort(key=lambda k: rank.get(k.tag.rsplit("}", 1)[-1], 999))
    for k in kids:
        el.remove(k)
    for k in kids:
        el.append(k)


def _tidy(root):
    for tag in ("rPr", "pPr", "tblPr", "tcPr"):
        for el in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + tag):
            _order(el)


def _run(par, s, rtl):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    text = _fix_ar(s["text"]) if rtl else s["text"]
    r = par.add_run(text)
    f = r.font
    f.size = Pt(round(s["size"] * 2) / 2)
    f.bold = s["bold"]
    f.italic = s["italic"]
    c = s["color"]
    if c:
        f.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
    rpr = r._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = _el("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), s["font"])
    half = str(int(round(s["size"] * 2)))
    rpr.append(_el("w:szCs", w_val=half))
    if s["bold"]:
        rpr.append(_el("w:bCs"))
    if s["italic"]:
        rpr.append(_el("w:iCs"))
    if RTL_RE.search(text):
        rpr.append(_el("w:rtl"))
    return r


def _para_format(par, rtl, size, align):
    from docx.oxml.ns import qn
    ppr = par._p.get_or_add_pPr()
    if rtl:
        ppr.append(_el("w:bidi"))
    ppr.append(_el("w:spacing", w_before=0, w_after=0, w_line=int(round(size * 1.18 * TW)), w_lineRule="exact"))
    jc = {"start": "left", "end": "right", "center": "center"}[align]
    # (Word, and LibreOffice copying it, read "left" in a right-to-left paragraph as its start: the right edge)
    ppr.append(_el("w:jc", w_val=jc))
    ind = ppr.find(qn("w:ind"))
    if ind is not None:
        ppr.remove(ind)


def _frame(doc, line, page_w):
    x0, y0, x1, y1 = line["bbox"]
    size = line["size"]
    pad = size * 0.9 + 6 + (x1 - x0) * 0.08    # Word's letters are rarely exactly as wide: room to spare
    w = (x1 - x0) + pad
    lh = size * 1.18
    y = y0 + ((y1 - y0) - lh) / 2
    if line["rtl"]:                             # the right edge stays put, the room goes to the left
        x = x1 + 1 - w
        if x < 0:
            w, x = x1 + 1, 0.0
    else:                                       # the left edge stays put, the room goes to the right
        x = max(0.0, x0 - 1)
        w = min(w, page_w - x)
    par = doc.add_paragraph()
    _para_format(par, line["rtl"], size, "start")
    ppr = par._p.get_or_add_pPr()
    fp = _el("w:framePr", w_w=int(w * TW), w_hSpace=0, w_wrap="around", w_vAnchor="page", w_hAnchor="page",
             w_x=int(x * TW), w_y=int(max(0, y) * TW))
    ppr.insert(0, fp)
    for s in line["runs"]:
        _run(par, s, line["rtl"])


def _table(doc, bbox, rows, spans, bg, scale):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    xs = sorted({round(c[0], 1) for r in rows for c in r if c} | {round(c[2], 1) for r in rows for c in r if c})
    cols = []
    for x in xs:
        if not cols or x - cols[-1] > 1.5:
            cols.append(x)
    # drop thin empty slivers between ruling lines that aren't real columns
    k = 0
    while k < len(cols) - 1:
        a, b = cols[k], cols[k + 1]
        if b - a < 14 and len(cols) > 2 and not any(a <= (s["bbox"][0] + s["bbox"][2]) / 2 <= b and bbox[1] <= s["bbox"][1] <= bbox[3]
                                                     and s["text"].strip() for s in spans):
            del cols[k + 1 if k + 1 < len(cols) - 1 else k]
            continue
        k += 1
    ncol = len(cols) - 1
    if ncol < 1:
        return
    tbl = doc.add_table(rows=len(rows), cols=ncol)
    tpr = tbl._tbl.tblPr
    tpr.append(_el("w:tblpPr", w_leftFromText=0, w_rightFromText=0, w_topFromText=0, w_bottomFromText=0,
                   w_vertAnchor="page", w_horzAnchor="page", w_tblpX=int(bbox[0] * TW), w_tblpY=int(bbox[1] * TW)))
    tpr.append(_el("w:tblOverlap", w_val="overlap"))
    tpr.append(_el("w:tblLayout", w_type="fixed"))
    borders = _el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el(f"w:{side}", w_val="single", w_sz=6, w_space=0, w_color="000000"))
    tpr.append(borders)
    mar = _el("w:tblCellMar")
    for side in ("left", "right"):
        mar.append(_el(f"w:{side}", w_w=40, w_type="dxa"))
    tpr.append(mar)
    grid = tbl._tbl.tblGrid
    for gc, (a, b) in zip(grid.findall(qn("w:gridCol")), zip(cols, cols[1:])):
        gc.set(qn("w:w"), str(int((b - a) * TW)))

    def col_of(x):
        return min(range(len(cols)), key=lambda i: abs(cols[i] - x))

    for ri, cells in enumerate(rows):
        real = [c for c in cells if c]
        if not real:
            continue
        top, bot = min(c[1] for c in real), max(c[3] for c in real)
        tr = tbl.rows[ri]
        tr.height = Pt(bot - top)
        trpr = tr._tr.get_or_add_trPr()
        for h in trpr.findall(qn("w:trHeight")):
            h.set(qn("w:hRule"), "atLeast")
        used = set()
        for c in real:
            a, b = col_of(c[0]), col_of(c[2])
            if b <= a or a in used:
                continue
            cell = tbl.cell(ri, a)
            if b - a > 1:
                cell = cell.merge(tbl.cell(ri, b - 1))
            used.update(range(a, b))
            tcpr = cell._tc.get_or_add_tcPr()
            tcpr.append(_el("w:vAlign", w_val="center"))
            col = _sample(bg, c, scale)
            if col:
                tcpr.append(_el("w:shd", w_val="clear", w_color="auto", w_fill=col))
            inner = [s for s in spans if _inside(s["bbox"], c)]
            lines = _lines(inner)
            par = cell.paragraphs[0]
            if not lines:
                _para_format(par, False, 8, "start")
                continue
            for k, ln in enumerate(lines):
                if k:
                    par = cell.add_paragraph()
                lx0, _, lx1, _ = ln["bbox"]
                left, right = lx0 - c[0], c[2] - lx1
                if abs(left - right) < 0.25 * (c[2] - c[0]) or (left > 3 and right > 3 and abs(left - right) < 8):
                    align = "center"
                elif ln["rtl"]:
                    align = "start" if right <= left else "end"
                else:
                    align = "start" if left <= right else "end"
                _para_format(par, ln["rtl"], ln["size"], align)
                for s in ln["runs"]:
                    _run(par, s, ln["rtl"])
    doc.add_paragraph()        # a floating table needs a paragraph after it


def _sample(bg, rect, scale):
    """The fill colour inside a cell (from the page without its text), or None when it's white."""
    if bg is None:
        return None
    x0, y0, x1, y1 = (int(v * scale) for v in rect)
    if x1 - x0 < 6 or y1 - y0 < 6:
        return None
    from collections import Counter
    pts = [(x, y) for x in range(x0 + 3, x1 - 3, max(1, (x1 - x0) // 9)) for y in range(y0 + 3, y1 - 3, max(1, (y1 - y0) // 5))]
    cols = Counter(bg.getpixel((x, y))[:3] for x, y in pts if 0 <= x < bg.width and 0 <= y < bg.height)
    if not cols:
        return None
    (r, g, b), _n = cols.most_common(1)[0]
    if r > 245 and g > 245 and b > 245:
        return None
    return f"{r:02X}{g:02X}{b:02X}"


def _background(src_doc, pno, spans, tables, dpi=150):
    """The page without its text (lines, boxes, logos, pictures), as a picture; None when the page is blank."""
    import pymupdf
    from PIL import Image
    tmp = pymupdf.open()
    tmp.insert_pdf(src_doc, from_page=pno, to_page=pno)
    pg = tmp[0]
    for s in spans:
        pg.add_redact_annot(pymupdf.Rect(s["bbox"]), fill=False)
    try:
        pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
    except TypeError:                             # older PyMuPDF
        pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE)
    scale = dpi / 72
    pix = pg.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    tmp.close()
    full = img.copy()                             # with table fills, for sampling cell colours
    if tables:
        from PIL import ImageDraw
        dr = ImageDraw.Draw(img)
        for bbox, _rows in tables:               # the real Word tables draw these themselves
            dr.rectangle([int(bbox[0] * scale) - 2, int(bbox[1] * scale) - 2,
                          int(bbox[2] * scale) + 2, int(bbox[3] * scale) + 2], fill=(255, 255, 255))
    lo, hi = img.convert("L").getextrema()
    blank = lo > 250
    return (None if blank else img), full, scale


def _anchor_picture(doc, png: bytes, w_pt, h_pt):
    """A picture behind the text, at the top-left corner of the page."""
    from docx.oxml import parse_xml
    from docx.shared import Pt
    par = doc.add_paragraph()
    _para_format(par, False, 1, "start")
    run = par.add_run()
    inline = run.add_picture(io.BytesIO(png), width=Pt(w_pt), height=Pt(h_pt))._inline
    emu_w, emu_h = inline.extent.cx, inline.extent.cy
    graphic = inline.graphic
    from docx.oxml.ns import nsdecls
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{emu_w}" cy="{emu_h}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="{1000 + len(doc.inline_shapes)}" name="Page background"/><wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


def convert(src, out, progress=None, cancelled=None):
    import pymupdf
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.shared import Pt
    pdf = pymupdf.open(str(src))
    doc = Document()
    body = doc.element.body
    for p in list(body):                          # start empty (keep the section settings)
        if not p.tag.endswith("sectPr"):
            body.remove(p)
    n = pdf.page_count
    for pno in range(n):
        if cancelled and cancelled():
            break
        page = pdf[pno]
        rect = page.rect
        sec = doc.sections[-1] if pno == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        sec.page_width, sec.page_height = Pt(rect.width), Pt(rect.height)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Pt(0))
        sec.header_distance = sec.footer_distance = Pt(0)
        spans = _spans(page)
        tables = _good_tables(page, spans)
        bg, full, scale = _background(pdf, pno, spans, tables)
        if bg is not None:
            buf = io.BytesIO()
            bg.save(buf, "PNG", optimize=True)
            _anchor_picture(doc, buf.getvalue(), rect.width, rect.height)
        free = [s for s in spans if not any(_inside(s["bbox"], tb) for tb, _r in tables)]
        for ln in _lines(free):
            _frame(doc, ln, rect.width)
        for tb, rows in tables:
            _table(doc, tb, rows, spans, full, scale)
        if progress:
            progress((pno + 1) / n)
    pdf.close()
    if not len([e for e in body if not e.tag.endswith("sectPr")]):
        doc.add_paragraph()
    _tidy(doc.element.body)
    doc.save(str(out))
    return Path(out)
